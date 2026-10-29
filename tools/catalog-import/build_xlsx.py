"""Build the admin import workbook.  python build_xlsx.py            -> pipes & fittings
                                     python build_xlsx.py electrical -> electrical (after normalize_cables.py)"""
import json, collections, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os
KIND = sys.argv[1] if len(sys.argv) > 1 else 'pipes'
P = 'out/electrical_' if KIND == 'electrical' else 'out/'
OUT = 'out/BuildMate_Electrical_Import.xlsx' if KIND == 'electrical' else 'out/BuildMate_Pipes_Fittings_Import.xlsx'
rows = json.load(open(P + 'catalog.json')); issues = json.load(open(P + 'issues.json'))
retired = json.load(open(P + 'retired.json')) if os.path.exists(P + 'retired.json') else []
HEAD = ["Code","Name","Category","SubCategory","Brand","Unit","HSN","GST %","MRP","List Price","Partner Discount %","Retail Discount %",
        "Reward Points / Unit","Reorder Level","Sequence","Barcode","Aliases (comma separated)","Name (Hindi)","Name (Gujarati)",
        "Active","Show On Web","Show On App","Description","Image URL","Variant Family","Variant Label","Specs","Purchase Discount %"]
wb = Workbook(); ws = wb.active; ws.title = "Products"
ws.append(HEAD)
# Retired codes go in switched off (Active = No) so the live store matches the sheets.
for r in rows + retired:
    ws.append([r['code'], r['name'], r['category'], r['subCategory'], r['brand'], r['unit'], r['hsn'], r['gst'], r['mrp'] or None, r['listPrice'] or None,
               r['partnerDiscount'], r['retailDiscount'], 0, r['reorderLevel'], r['sequence'], None, ', '.join(r['aliases']), None, None,
               'Yes' if r['active'] else 'No', 'Yes', 'Yes', r['description'], r['image'], r['variantFamily'], r['variantLabel'],
               '; '.join(f'{k}: {v}' for k, v in r['specs'].items()), r['purchaseDiscount']])
bold = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='0F4FA8')
for c in ws[1]: c.font = bold; c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical='center')
widths = {'A':14,'B':52,'C':16,'D':15,'E':10,'W':50,'X':40,'Y':36,'Z':22,'AA':44}
for i in range(1, len(HEAD)+1):
    L = get_column_letter(i); ws.column_dimensions[L].width = widths.get(L, 12)
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions

# Needs checking: per-row problems, with blank purchase discounts rolled up per family
ck = wb.create_sheet("Needs checking")
ck.append(["Code", "Product", "Where", "Problem", "What I did"])
fam_of = {r['code']: r['variantFamily'] for r in rows}
blank = collections.defaultdict(list)
for i in issues:
    if i['problem'].startswith('Purchase discount blank'): blank[fam_of.get(i['code'], '?')].append(i); continue
    ck.append([i['code'], i['name'], i['source'], i['problem'], i['action']])
if os.path.exists(P + 'images_report.json'):
    rep = json.load(open(P + 'images_report.json'))
    counts = collections.Counter(r['variantFamily'] for r in rows if not r['image'])
    for fam in rep['familiesWithoutImage']:
        ck.append([f"{counts[fam]} codes", fam, "image sheets", "No product image for this " + ("line" if KIND == 'electrical' else "range"), "Shows a placeholder — add a row to the image sheet"])
    for x in rep['unusedImageLinks']:
        ck.append(["—", f"{x['brand']} {x['family']} {x['type']} {x['class']}".strip(), x['sheet'], "Image link matched no product (check spelling of item / class)", "Not used"])
for fam, items in blank.items():
    ck.append([f"{len(items)} codes", fam, items[0]['source'].split(' row')[0], "Purchase discount blank for this whole range", "Cost left empty — fill in when known"])
for c in ck[1]: c.font = bold; c.fill = fill
for L, w in zip('ABCDE', (16, 52, 46, 70, 44)): ck.column_dimensions[L].width = w
ck.freeze_panes = 'A2'

if retired:
    rm = wb.create_sheet("Removed")
    rm.append(["Code", "Product", "Product family", "What happens on import"])
    for r in retired: rm.append([r['code'], r['name'], r['variantFamily'], "Switched off (hidden from website and app). Kept for past orders and invoices."])
    for c in rm[1]: c.font = bold; c.fill = fill
    for L, w in zip('ABCD', (16, 56, 36, 64)): rm.column_dimensions[L].width = w

rn = wb.create_sheet("Renamed")
rn.append(["In your sheet", "Shown in the store as", "Why"])
RENAMED = json.load(open(P + 'renamed.json')) if os.path.exists(P + 'renamed.json') else [("M.T.A. / F.T.A.", "MTA / FTA", "Same spelling across brands"), ("R BUSH", "Reducer Bush", "Same name as UPVC/CPVC"),
                  ("TRANSACTION BUSH", "Transition Bush", "Typo"), ("MUULTIFLOOR TRAP", "Multifloor Trap", "Typo"),
                  ("P.T.F.E. SEAL TAP (Teflon)", "PTFE Seal Tape (Teflon)", "Typo"), ("(Hexagone)", "(HEX)", "Typo"),
                  ("4 KG / 6 KG / 10 KG", "4 / 6 / 10 kg/cm²", "Pressure class unit"),
                  ("Sizes: 15 MM (½\") / 15mm (½\") / 20X15 / 25x15 / 25×15mm", "15 mm (½\") / 20×15 mm (¾\" × ½\")", "One format everywhere"),
                  ("Product names", "Brand + Material + Item + Class – Size", "Consistent names so search and the size picker work")]
for a, b, why in RENAMED:
    rn.append([a, b, why])
for c in rn[1]: c.font = bold; c.fill = fill
for L, w in zip('ABC', (48, 44, 44)): rn.column_dimensions[L].width = w
wb.save(OUT)
print('rows', ws.max_row - 1, '(of which switched off as removed:', len(retired), ') checks', ck.max_row - 1)
