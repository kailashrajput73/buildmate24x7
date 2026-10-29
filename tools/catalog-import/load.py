import openpyxl,re
import os, sys
# Folder holding the four brand sheets (pass as first argument or SHEETS_DIR).
D=os.path.join(sys.argv[1] if len(sys.argv) > 1 else os.environ.get('SHEETS_DIR', 'sheets'), '')
def norm(h): return re.sub(r'[^a-z]','',str(h or '').lower())
ALIAS={'category':'category','type':'type','subcategory':'subcat','class':'cls','brand':'brand','hsncode':'hsn','gst':'gst','uom':'uom',
 'productfamily':'family','product':'family','productname':'name','sizemmcmin':'size','sizecm':'size','length':'length','productcode':'code',
 'mrprspernos':'mrp','discount':'disc','sellingprice':'sell','packsize':'pack','mrppkg':'mrppkg','rol':'rol','status':'status','imageurl':'image','displaysequence':'seq'}
def sheet(f,name=None):
    wb=openpyxl.load_workbook(D+f,read_only=True,data_only=True)
    ws=wb[name] if name and name in wb.sheetnames else wb.worksheets[0]
    name=ws.title
    allr=list(ws.iter_rows(values_only=True)); hdr=allr[0]
    body=[(n,r) for n,r in enumerate(allr[1:],start=2) if any(c not in (None,'') for c in r)]
    # A header can appear twice (e.g. two "Status" columns): use the one that actually has values.
    cand={}
    for i,h in enumerate(hdr):
        k=ALIAS.get(norm(h))
        if k: cand.setdefault(k,[]).append(i)
    cols={k:max(v,key=lambda i: sum(1 for _,r in body if i<len(r) and r[i] not in (None,''))) for k,v in cand.items()}
    out=[]
    for n,r in body:
        d={k:(r[i] if i<len(r) else None) for k,i in cols.items()}
        d=  {k:(v.strip() if isinstance(v,str) else v) for k,v in d.items()}
        d['_row']=n; d['_src']=f+':'+name; out.append(d)
    return out
